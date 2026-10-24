"""Compare English and Kiswahili songs in the database with their PowerPoint sources (python-pptx, independent
of the app's own reader). Slide order, every line and every line break are compared."""
import sqlite3, os, re, sys, json, unicodedata
from pptx import Presentation

DB = os.environ.get('MF_DB', r"E:\My Projects\MessageFlow\database\messageflow.db")
con = sqlite3.connect(f'file:{DB}?mode=ro', uri=True)


def ws(s):
    s = unicodedata.normalize('NFC', s).replace('\u00a0', ' ').replace('\u200b', '')
    return re.sub(r'[ \t]+', ' ', s).strip()


def shape_lines(shape):
    out = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            # python-pptx renders <a:br/> as \x0b inside paragraph text
            for piece in p.text.split('\x0b'):
                out.append(piece)
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            out += shape_lines(s)
    if getattr(shape, 'has_table', False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    out += p.text.split('\x0b')
    return out


def slide_lines(path):
    prs = Presentation(path)
    slides = []
    for sl in prs.slides:
        lines = []
        for sh in sl.shapes:
            lines += shape_lines(sh)
        slides.append([ws(l) for l in lines if ws(l)])
    return slides


def main(lang, root, out):
    songs = con.execute('select Id, Title, SourceFilePath, FileName from Songs where Language=? and IsActive=1 order by Id',
                        (lang,)).fetchall()
    results = []
    for sid, title, sp, fn in songs:
        path = re.sub(r'^[Dd]:', 'E:', sp)
        if not os.path.exists(path):
            path = os.path.join(root, fn)
        secs = con.execute('select SectionOrder, SectionLabel, Text from SongSections where SongId=? order by SectionOrder',
                           (sid,)).fetchall()
        r = {'id': sid, 'title': title, 'file': fn, 'sections': len(secs), 'issues': []}
        if not os.path.exists(path):
            r['issues'].append({'kind': 'NO_SOURCE', 'detail': path})
            results.append(r); continue
        try:
            slides = slide_lines(path)
        except Exception as e:
            r['issues'].append({'kind': 'READ_ERROR', 'detail': str(e)})
            results.append(r); continue
        nonempty = [s for s in slides if s]
        r['slides'] = len(nonempty)
        dbsecs = [[ws(l) for l in t.replace('\r\n', '\n').split('\n') if ws(l)] for _, _, t in secs]
        # 1. every database line, in order, must be a source line in order (nothing invented, nothing reordered)
        flat_db = [(i, l) for i, s in enumerate(dbsecs) for l in s]
        flat_src = [(i, l) for i, s in enumerate(nonempty) for l in s]
        if [l for _, l in flat_db] != [l for _, l in flat_src]:
            # find first mismatch
            a = [l for _, l in flat_db]; b = [l for _, l in flat_src]
            k = next((i for i in range(min(len(a), len(b))) if a[i] != b[i]), min(len(a), len(b)))
            r['issues'].append({'kind': 'LINES_DIFFER', 'db_lines': len(a), 'src_lines': len(b), 'at': k,
                                'db': a[k:k + 2], 'src': b[k:k + 2]})
        # 2. section boundaries = slide boundaries
        elif [len(s) for s in dbsecs] != [len(s) for s in nonempty]:
            r['issues'].append({'kind': 'SECTION_SPLIT_DIFFERS', 'db': [len(s) for s in dbsecs],
                                'src': [len(s) for s in nonempty]})
        for i, (_, lab, t) in enumerate(secs):
            if not t.strip():
                r['issues'].append({'kind': 'EMPTY_SECTION', 'section': i + 1})
        results.append(r)
    with open(out, 'w', encoding='utf-8') as f:
        for r in results:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')
    bad = [r for r in results if r['issues']]
    print(f'{lang}: songs={len(results)} identical={len(results) - len(bad)} with_issues={len(bad)}')
    import collections
    print(collections.Counter(i['kind'] for r in bad for i in r['issues']))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3])
